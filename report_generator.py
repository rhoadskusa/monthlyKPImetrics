"""
Build the monthly summary table and write it to Excel + CSV.
"""

from __future__ import annotations

import logging
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

COLUMNS = ["Metric", "Actual", "Budget", "Variance", "Last Month", "YTD", "Same Period LY"]


def build_summary_table(rows: list[dict]) -> pd.DataFrame:
    return pd.DataFrame(rows, columns=COLUMNS)


def write_reports(df: pd.DataFrame, period: date, output_dir: Path) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = f"{period.year}_{period.month:02d}_metrics_report"

    csv_path = output_dir / f"{stem}.csv"
    df.to_csv(csv_path, index=False)
    log.info("Wrote %s", csv_path)

    xlsx_path = output_dir / f"{stem}.xlsx"
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Summary", index=False)
        ws = writer.sheets["Summary"]
        for col_idx, column in enumerate(df.columns, start=1):
            ws.cell(row=1, column=col_idx).font = Font(bold=True)
            max_len = max(len(str(column)), df[column].astype(str).map(len).max() if len(df) else 0)
            ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 2
    log.info("Wrote %s", xlsx_path)

    return csv_path, xlsx_path
